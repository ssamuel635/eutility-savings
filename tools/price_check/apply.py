#!/usr/bin/env python3
"""
Weekly price check: step 3 (apply approved changes to the price sheet).

Reads price-checks/<date>/changes.json and writes the approved plan and promotion
changes into data/price-comparison-sheet.xlsx. Blocked items (failed an accuracy
check) are never applied. "Couldn't verify" plans are left exactly as they are.

What counts as approved comes from approved-<date>.json, saved by the "Save
approvals" button in report.html. It is looked for in the run folder, then in
Downloads, and a copy is kept in the run folder as the record.

Usage:
  python3 tools/price_check/apply.py price-checks/2026-10-12                  # what the approver ticked
  python3 tools/price_check/apply.py price-checks/2026-10-12 --approved FILE  # a specific approvals file
  python3 tools/price_check/apply.py price-checks/2026-10-12 --dry-run        # show what would change, save nothing
  Without an approvals file (by hand, approval given some other way):
  python3 tools/price_check/apply.py price-checks/2026-10-12 --all            # everything not blocked
  python3 tools/price_check/apply.py price-checks/2026-10-12 --only id1,id2   # just these (--skip to leave some out)
  python3 tools/price_check/apply.py ... --no-promos                          # prices only

Then: python3 tools/build_data.py --rates-date <date>  &&  npm test
"""
import argparse
import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

from runlog import record

ROOT = Path(__file__).resolve().parents[2]
SHEET = ROOT / "data" / "price-comparison-sheet.xlsx"


def slug(*parts):
    return re.sub(r"[^a-z0-9]+", "-", "-".join(str(p) for p in parts).lower()).strip("-")


def is_off_peak(window):  # same rule as tools/build_data.py
    w = (window or "").lower()
    if "night" in w or "weekend" in w:
        return True
    m = re.match(r"\s*(\d{1,2}):", w)
    return bool(m and int(m.group(1)) >= 18)


def tier_start(label):
    s = str(label).replace("–", "-")
    m = re.match(r"\s*>\s*(\d+)", s) or re.match(r"\s*(\d+)", s)
    return float(m.group(1)) if m else 0.0


def fmt_date(iso):
    if not iso:
        return "—"
    d = date.fromisoformat(iso)
    return f"{d.day} {d.strftime('%b %Y')}"


def find_approvals(run, given):
    """The approvals file for this run: the one given, else the run folder, else Downloads (newest)."""
    if given:
        return Path(given)
    for name in (f"approved-{run.name}.json", "approved.json"):
        if (run / name).exists():
            return run / name
    hits = sorted((Path.home() / "Downloads").glob(f"approved-{run.name}*.json"), key=lambda f: f.stat().st_mtime)
    return hits[-1] if hits else None


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # Windows consoles default to cp1252
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("run_dir")
    ap.add_argument("--only", default="")
    ap.add_argument("--skip", default="")
    ap.add_argument("--no-promos", action="store_true")
    ap.add_argument("--approved", default="", help="approvals file (default: approved-<date>.json in the run folder or Downloads)")
    ap.add_argument("--all", action="store_true", help="no approvals file: apply everything not blocked")
    ap.add_argument("--dry-run", action="store_true", help="print what would change; save nothing")
    a = ap.parse_args()
    run = (ROOT / a.run_dir) if not Path(a.run_dir).is_absolute() else Path(a.run_dir)
    ch = json.loads((run / "changes.json").read_text(encoding="utf-8"))
    only = {s for s in a.only.split(",") if s}
    skip = {s for s in a.skip.split(",") if s}

    approval = None
    by_hand = a.all or only or skip
    src = None if (by_hand and not a.approved) else find_approvals(run, a.approved)
    if src:
        approval = json.loads(src.read_text(encoding="utf-8"))
        if approval.get("run") != run.name:
            sys.exit(f"{src.name} is for the {approval.get('run')} run, not {run.name}. Nothing applied.")
        print(f"Approved by {approval.get('approved_by')} at {approval.get('approved_at')} ({src})")
    elif not by_hand:
        sys.exit(f"No approvals found for {run.name}. Open {run / 'report.html'}, tick what's right and press "
                 "Save approvals (or use --all / --only if approval was given another way). Nothing applied.")
    plans_ok = set(approval["plans"]) if approval else None
    promos_ok = set(approval["promotions"]) if approval else None
    in_scope = lambda pid: pid and (not only or pid in only) and pid not in skip
    ok = lambda pid: in_scope(pid) and (plans_ok is None or pid in plans_ok)
    promo_ok = lambda pid: in_scope(pid) and (promos_ok is None or pid in promos_ok)

    wb = openpyxl.load_workbook(SHEET)
    ws = wb["Plans"]
    rows_by_id = {}
    for i in range(2, ws.max_row + 1):
        r, n = ws.cell(i, 1).value, ws.cell(i, 2).value
        if r and n and not str(r).startswith("*"):
            rows_by_id.setdefault(slug(r, n), []).append(i)

    log, delete = [], []
    for x in ch["plans"]:
        pid, t = x.get("id"), x.get("trust") or {}
        if not ok(pid) or t.get("blocked") or x["kind"] not in ("changed", "removed"):
            if x["kind"] in ("changed", "removed") and t.get("blocked"):
                log.append(f"BLOCKED, not applied: {x['plan']} ({'; '.join(t.get('flags', []) + t.get('disagreements', []))})")
            continue
        rows = rows_by_id.get(pid)
        if not rows:
            log.append(f"not in sheet, skipped: {pid}")
            continue
        if x["kind"] == "removed":
            delete += rows
            log.append(f"removed: {x['plan']}")
            continue
        p = x["proposed"]
        kind = str(ws.cell(rows[0], 4).value or "").lower()
        for i in rows:
            if p.get("contractMonths") is not None:
                ws.cell(i, 5).value = p["contractMonths"]
            if p.get("dailyCharge") is not None and ws.cell(i, 9).value not in ("—",):
                ws.cell(i, 9).value = p["dailyCharge"]
        if "time of use" in kind and p.get("tou"):
            for i in rows:
                band = "offPeak" if is_off_peak(ws.cell(i, 6).value) else "peak"
                ws.cell(i, 7).value = p["tou"][band]["rate"]
        elif "tiered" in kind and p.get("tiers"):
            for i, b in zip(sorted(rows, key=lambda i: tier_start(ws.cell(i, 6).value)), p["tiers"]):
                ws.cell(i, 7).value = b["rate"]
        elif "discount" in kind and p.get("discount"):
            d = p["discount"]
            v = f"{d['value']:g}"
            ws.cell(rows[0], 8).value = f"{v}*" if d["kind"] == "cents" else v
        elif p.get("rate") is not None:
            ws.cell(rows[0], 7).value = p["rate"]
        log.append(f"updated: {x['plan']} — " + "; ".join(f"{f}: {o} → {n}" for f, o, n in x["changes"]))
    for i in sorted(set(delete), reverse=True):
        ws.delete_rows(i)

    if not a.no_promos:
        pw = wb["Promotions"]
        prow = {}
        for i in range(2, pw.max_row + 1):
            r, n = pw.cell(i, 1).value, pw.cell(i, 2).value
            if r and n:
                prow[slug(r, n)] = i
        names = {x["id"]: (x["retailer"], x["plan"]) for x in ch["plans"] if x.get("id")}
        removed = {x["id"] for x in ch["plans"] if x["kind"] == "removed" and ok(x["id"]) and not (x.get("trust") or {}).get("blocked")}
        pdel = [prow[pid] for pid in removed if pid in prow]
        last = max([i for i in prow.values()] or [1])
        for pr in ch["promotions"]:
            pid = pr["id"]
            if not promo_ok(pid) or pid in removed:
                continue
            new = pr["new"]
            if not new:
                if pid in prow:
                    pdel.append(prow[pid])
                    log.append(f"promotion ended: {pr['plan']}")
                continue
            i = prow.get(pid)
            if not i:
                last += 1
                pw.insert_rows(last)
                i = prow[pid] = last
                pw.cell(i, 1).value, pw.cell(i, 2).value = names.get(pid, (pr["retailer"], pr["plan"]))
            vals = [new.get("title") or "—", new.get("code") or "—", new.get("conditions") or "—",
                    (new.get("value") or "—").replace("S$", "$"), "—", fmt_date(new.get("expires")), pr["retailer"],
                    fmt_date(run.name)]
            for c, v in enumerate(vals, start=3):
                pw.cell(i, c).value = v
            log.append(f"promotion: {pr['plan']} — {new.get('title')}")
        for i in sorted(set(pdel), reverse=True):
            pw.delete_rows(i)

    print("\n".join(log) or "Nothing to apply.")
    if a.dry_run:
        print("\nDry run: nothing saved.")
        return
    wb.save(SHEET)
    if src and src.parent.resolve() != run.resolve():
        (run / f"approved-{run.name}.json").write_text(src.read_text(encoding="utf-8"), encoding="utf-8")
    record(run.parent, run.name,
           approved_by=approval.get("approved_by") if approval else "by hand (--all/--only)",
           approved_at=approval.get("approved_at") if approval else "",
           plans_approved=sum(1 for l in log if l.startswith(("updated:", "removed:"))),
           promos_approved=sum(1 for l in log if l.startswith(("promotion:", "promotion ended:"))),
           applied_at=date.today().isoformat())
    print(f"\nSaved {SHEET.relative_to(ROOT)}. Next: python3 tools/build_data.py --rates-date {run.name} && npm test")


if __name__ == "__main__":
    sys.exit(main())
