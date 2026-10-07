#!/usr/bin/env python3
"""
Convert the Eutility price comparison sheet (Excel) into the JS data module
used by the site: assets/js/data/plans.js

Usage:
    pip install openpyxl
    python3 tools/build_data.py                       # uses data/price-comparison-sheet.xlsx
    python3 tools/build_data.py path/to/sheet.xlsx    # or a newer copy

Sheet contract (same as the file supplied Sep 2026):
  Plans tab       : Retailer | Plan Name | Product Category | Pricing Type | Contract (Months) |
                    Time / Usage | Rate (¢/kWh) | Discount (%) | Daily Charge ($/day) |
                    Security Deposit / Requirement | ETF ($) | Source
  Promotions tab  : Retailer | Plan Name | Promotion | Promo Code | Eligibility / Conditions |
                    Promo Value | Start Date | Expiry Date | Source | Last Updated
Rows sharing Retailer + Plan Name are merged (Time-of-Use and Tiered plans use one row per band).
"""
import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data" / "price-comparison-sheet.xlsx"
OUT = ROOT / "assets" / "js" / "data" / "plans.js"

DASHES = {"—", "-", "–", "", None}


def clean(v):
    if isinstance(v, str):
        v = v.strip()
    return None if v in DASHES else v


def num(v):
    v = clean(v)
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"-?\d+(\.\d+)?", str(v))
    return float(m.group()) if m else None


def slug(*parts):
    s = "-".join(str(p) for p in parts).lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def pricing_kind(ptype):
    p = (ptype or "").lower()
    if "regulated" in p:
        return "regulated"
    if "time of use" in p:
        return "tou"
    if "tiered" in p:
        return "tiered"
    if "discount" in p:
        return "dot"
    if "fixed" in p:
        return "fixed"
    return "unknown"


def parse_tier(label):
    """'0–300 kWh' -> (0, 300); '301–600 kWh' -> (300, 600); '>600 kWh' -> (600, None)"""
    s = str(label).replace("–", "-")
    m = re.match(r"\s*>\s*(\d+)", s)
    if m:
        return float(m.group(1)), None
    m = re.match(r"\s*(\d+)\s*-\s*(\d+)", s)
    if m:
        lo, hi = float(m.group(1)), float(m.group(2))
        return (lo - 1 if lo > 0 else 0.0), hi
    return None, None


def is_off_peak(window):
    w = (window or "").lower()
    if "night" in w or "weekend" in w:
        return True
    # windows that start in the evening (19:00, 21:00, 23:00) are off-peak
    m = re.match(r"\s*(\d{1,2}):", w)
    return bool(m and int(m.group(1)) >= 18)


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Plans"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    plans, order = {}, []
    for r in rows:
        retailer, name = clean(r[0]), clean(r[1])
        if not retailer or not name or str(retailer).startswith("*"):
            continue
        if len(r) < 12 or (r[3] is None and r[6] is None and r[10] is None):
            continue
        key = (retailer, name)
        kind = pricing_kind(clean(r[3]))
        contract = num(r[4])
        window = clean(r[5])
        rate = num(r[6])
        disc_raw = clean(r[7])
        daily = num(r[8]) or 0.0

        if key not in plans:
            available = kind != "unknown" and "not yet" not in str(name).lower()
            p = {
                "id": slug(retailer, name),
                "retailer": retailer,
                "name": name,
                "category": clean(r[2]) or "Standard",
                "pricingType": kind,
                "pricingLabel": clean(r[3]),
                "contractMonths": int(contract) if contract is not None else None,
                "rate": None,
                "discount": None,
                "dailyCharge": daily,
                "tou": None,
                "tiers": None,
                "deposit": clean(r[9]),
                "etf": clean(r[10]) if not isinstance(r[10], (int, float)) else (
                    "$0" if r[10] == 0 else f"${r[10]}"),
                "source": clean(r[11]),
                "available": available,
                "promo": None,
                "notes": [],
            }
            plans[key] = p
            order.append(key)
        p = plans[key]

        if kind in ("fixed", "regulated"):
            p["rate"] = rate
        elif kind == "dot":
            s = str(disc_raw or "")
            p["discount"] = {
                "value": num(s),
                "kind": "cents" if "*" in s else "percent",
            }
        elif kind == "tou":
            band = {"window": window, "rate": rate}
            p["tou"] = p["tou"] or {}
            p["tou"]["offPeak" if is_off_peak(window) else "peak"] = band
            if window and "weekend" in window.lower():
                p["tou"]["weekendOffPeak"] = True
        elif kind == "tiered":
            lo, hi = parse_tier(window)
            p["tiers"] = (p["tiers"] or []) + [{"from": lo, "to": hi, "rate": rate}]

    # Promotions
    promos = {}
    if "Promotions" in wb.sheetnames:
        for r in wb["Promotions"].iter_rows(min_row=2, values_only=True):
            retailer, name = clean(r[0]), clean(r[1])
            if not retailer or not name or (retailer, name) not in plans:
                continue
            promos[(retailer, name)] = {
                "title": clean(r[2]),
                "code": clean(r[3]),
                "conditions": clean(r[4]),
                "value": clean(r[5]),
                "expires": clean(r[7]) if not hasattr(r[7], "isoformat") else r[7].isoformat(),
            }
    for key, pr in promos.items():
        plans[key]["promo"] = pr
        cond = (pr.get("conditions") or "").lower()
        if "referral-only" in cond or "requires a referral code" in cond:
            plans[key]["referralOnly"] = True
            plans[key]["notes"].append("Referral-only plan (requires a referral code).")

    out = []
    for key in order:
        p = plans[key]
        if p["tiers"]:
            p["tiers"].sort(key=lambda t: t["from"] or 0)
        if p["pricingType"] == "dot":
            p["notes"].append("Discount off SP tariff — price moves when the tariff is revised quarterly.")
        if p["pricingType"] == "tou":
            p["notes"].append("Time-of-use — estimate depends on how much you use in off-peak hours.")
        out.append(p)

    tariff = next((p["rate"] for p in out if p["pricingType"] == "regulated"), None)
    meta = {
        "source": SRC.name,
        "generated": date.today().isoformat(),
        "regulatedTariff": tariff,
        "planCount": sum(1 for p in out if p["available"]),
    }

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        "// AUTO-GENERATED by tools/build_data.py — do not edit by hand.\n"
        f"// Source: {SRC.name}\n"
        f"export const META = {json.dumps(meta, ensure_ascii=False, indent=2)};\n\n"
        f"export const PLANS = {json.dumps(out, ensure_ascii=False, indent=2)};\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUT.relative_to(ROOT)} — {meta['planCount']} available plans, tariff {tariff}¢/kWh")


if __name__ == "__main__":
    main()
